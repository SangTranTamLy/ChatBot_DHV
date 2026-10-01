"""Import the user-provided canonical Phase 5 workbook without generating Q&A."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "data" / "evaluation" / "raw" / "Nhóm 4_Thu_Thap_Q&A.xlsx"
DEFAULT_MASTER = ROOT / "data" / "evaluation" / "qa_master_500.jsonl"
DEFAULT_LOCK = ROOT / "data" / "evaluation" / "qa_master_500.lock"
DEFAULT_AUDIT = ROOT / "data" / "evaluation" / "qa_master_500.audit.json"
EXPECTED_ROWS = 500
HEADERS = ("TT", "Câu hỏi", "Trả lời", "Nguồn")


def _text(value: object) -> str:
    """Convert a cell to text without rewriting its content."""

    return "" if value is None else str(value)


def _normalise(value: object) -> str:
    text = unicodedata.normalize("NFKC", _text(value)).casefold()
    return re.sub(r"\s+", " ", text).strip()


def _row_hash(values: tuple[str, str, str, str]) -> str:
    payload = json.dumps(values, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _load_workbook_rows(source_path: Path) -> tuple[list[dict[str, object]], dict[str, object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency failure
        raise RuntimeError(
            "Phase 5 Excel import requires openpyxl. "
            "Run: pip install -r requirements.txt"
        ) from exc

    if not source_path.exists():
        raise FileNotFoundError(f"canonical workbook not found: {source_path}")

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("canonical workbook must contain exactly one worksheet")
        worksheet = workbook.worksheets[0]
        header_row_number: int | None = None
        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=min(worksheet.max_row, 20),
                max_col=4,
                values_only=True,
            ),
            start=1,
        ):
            if tuple(_text(value).strip() for value in values[:4]) == HEADERS:
                header_row_number = row_number
                break
        if header_row_number is None:
            raise ValueError(f"missing required header row {HEADERS!r}")

        rows: list[dict[str, object]] = []
        for excel_row, values in enumerate(
            worksheet.iter_rows(min_row=header_row_number + 1, max_col=4, values_only=True),
            start=header_row_number + 1,
        ):
            cells = tuple(_text(value) for value in values[:4])
            if not any(cell.strip() for cell in cells):
                continue
            tt, question, answer, source = cells
            if not all(cell.strip() for cell in cells):
                raise ValueError(f"row {excel_row} has a blank TT, question, answer, or source")
            try:
                tt_number = int(tt.strip())
            except ValueError as exc:
                raise ValueError(f"row {excel_row} has a non-numeric TT: {tt!r}") from exc
            rows.append(
                {
                    "id": f"UQ{tt_number:04d}",
                    "tt": tt,
                    "question": question,
                    "reference_answer": answer,
                    "source": source,
                    "dataset_source_file": source_path.name,
                    "source_excel_row": excel_row,
                    "source_row_sha256": _row_hash(cells),
                    "split": "canonical",
                }
            )
        metadata = {
            "source_file": source_path.name,
            "source_path": str(source_path),
            "worksheet": worksheet.title,
            "header_row": header_row_number,
            "physical_rows": worksheet.max_row,
        }
        return rows, metadata
    finally:
        workbook.close()


def audit_rows(rows: list[dict[str, object]]) -> dict[str, object]:
    ids = Counter(str(row.get("id") or "") for row in rows)
    questions = Counter(_normalise(row.get("question")) for row in rows)
    answers = Counter(_normalise(row.get("reference_answer")) for row in rows)
    expected_ids = [f"UQ{index:04d}" for index in range(1, EXPECTED_ROWS + 1)]
    actual_ids = [str(row.get("id") or "") for row in rows]
    return {
        "total": len(rows),
        "expected_total": EXPECTED_ROWS,
        "canonical_only": all(row.get("split") == "canonical" for row in rows),
        "unique_ids": len(ids) == len(rows) and "" not in ids,
        "duplicate_ids": {key: value for key, value in ids.items() if value > 1},
        "duplicate_questions": {key: value for key, value in questions.items() if key and value > 1},
        "duplicate_answers": {key: value for key, value in answers.items() if key and value > 1},
        "missing_questions": [str(row.get("id")) for row in rows if not str(row.get("question") or "").strip()],
        "missing_answers": [str(row.get("id")) for row in rows if not str(row.get("reference_answer") or "").strip()],
        "missing_sources": [str(row.get("id")) for row in rows if not str(row.get("source") or "").strip()],
        "sequential_tt": actual_ids == expected_ids,
        "valid": (
            len(rows) == EXPECTED_ROWS
            and actual_ids == expected_ids
            and len(ids) == EXPECTED_ROWS
            and not any(key for key, value in questions.items() if key and value > 1)
            and not any(key for key, value in answers.items() if key and value > 1)
            and not any(row.get("split") != "canonical" for row in rows)
            and not any(
                not str(row.get(field) or "").strip()
                for row in rows
                for field in ("question", "reference_answer", "source")
            )
        ),
    }


def write_jsonl(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "".join(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n" for row in rows),
        encoding="utf-8",
    )


def build_dataset(
    *,
    source_path: Path = DEFAULT_SOURCE,
    master_path: Path = DEFAULT_MASTER,
    lock_path: Path = DEFAULT_LOCK,
    audit_path: Path = DEFAULT_AUDIT,
) -> dict[str, object]:
    rows, source_metadata = _load_workbook_rows(source_path)
    audit = audit_rows(rows)
    if not audit["valid"]:
        raise ValueError(
            f"canonical dataset validation failed: {json.dumps(audit, ensure_ascii=False)}"
        )

    write_jsonl(master_path, rows)
    dataset_hash = hashlib.sha256(master_path.read_bytes()).hexdigest()
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    lock_path.write_text(dataset_hash + "\n", encoding="utf-8")
    result: dict[str, object] = {
        "dataset": "user_canonical_500",
        "source": source_metadata,
        "audit": audit,
        "master_path": str(master_path),
        "master_sha256": dataset_hash,
        "lock_path": str(lock_path),
        "generated_questions": 0,
        "generated_answers": 0,
    }
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    audit_path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--master", type=Path, default=DEFAULT_MASTER)
    parser.add_argument("--lock", type=Path, default=DEFAULT_LOCK)
    parser.add_argument("--audit", type=Path, default=DEFAULT_AUDIT)
    args = parser.parse_args()
    result = build_dataset(
        source_path=args.source,
        master_path=args.master,
        lock_path=args.lock,
        audit_path=args.audit,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()


__all__ = ["audit_rows", "build_dataset", "main"]
