"""Convert a 2026 admission Q&A workbook into verified Structured JSON."""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.config.settings import PROJECT_ROOT

from .structured_json import SCHOOL_NAME, write_structured_json


DEFAULT_INPUT = PROJECT_ROOT / "data" / "raw" / "qa_tuyen_sinh" / "Nhom_4_Thu_Thap_QA.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "data" / "processed" / "qa_tuyen_sinh" / "qa_tuyen_sinh_2026.json"
SOURCE_URL = "https://dhv.edu.vn"


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _column_indexes(headers: tuple[str, ...]) -> tuple[int, int, int | None]:
    normalized = {header.casefold(): index for index, header in enumerate(headers) if header}
    question = next((normalized[key] for key in normalized if "câu hỏi" in key or "cau hoi" in key), None)
    answer = next((normalized[key] for key in normalized if "trả lời" in key or "tra loi" in key), None)
    source = next((normalized[key] for key in normalized if "nguồn" in key or "nguon" in key), None)
    if question is None or answer is None:
        raise ValueError("workbook must contain Câu hỏi and Trả lời columns")
    return question, answer, source


def convert_workbook(input_path: str | Path, output_path: str | Path) -> Path:
    """Convert non-empty Q&A rows while preserving workbook row provenance."""

    workbook_path = Path(input_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if not workbook.worksheets:
        raise ValueError("workbook contains no worksheets")
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    header_row = next((index for index, row in enumerate(rows) if any(_text(value) == "Câu hỏi" for value in row)), None)
    if header_row is None:
        raise ValueError("workbook header row was not found")
    question_index, answer_index, source_index = _column_indexes(tuple(_text(value) for value in rows[header_row]))

    records: list[dict[str, Any]] = []
    page_text: list[str] = []
    for row_number, row in enumerate(rows[header_row + 1 :], start=header_row + 2):
        question = _text(row[question_index] if question_index < len(row) else None)
        answer = _text(row[answer_index] if answer_index < len(row) else None)
        if not question and not answer:
            continue
        if not question or not answer:
            raise ValueError(f"row {row_number} must contain both question and answer")
        source = _text(row[source_index] if source_index is not None and source_index < len(row) else None)
        record_id = f"qa_{len(records) + 1:04d}"
        records.append(
            {
                "record_type": "qa_pair",
                "record_id": record_id,
                "question": question,
                "answer": answer,
                "text": f"Câu hỏi: {question}\nTrả lời: {answer}",
                "source_note": source,
                "workbook_row": row_number,
                "page": 1,
            }
        )
        page_text.append(f"Câu hỏi: {question}\nTrả lời: {answer}")

    if not records:
        raise ValueError("workbook contains no Q&A rows")
    text = "\n\n".join(page_text)
    document = {
        "document_id": "qa_tuyen_sinh_2026",
        "title": "Bộ câu hỏi và trả lời tư vấn tuyển sinh DHV 2026",
        "category": "qa_tuyen_sinh",
        "year": 2026,
        "data_role": "qa",
        "source": {
            "file_name": workbook_path.name,
            "raw_file": workbook_path.as_posix(),
            "source_url": SOURCE_URL,
            "source_urls": [SOURCE_URL],
            "organization": SCHOOL_NAME,
            "verified": True,
            "status": "verified",
            "verification_status": "verified",
            "source_date": "2026",
            "collected_at": "2026-09-26",
            "school_code": "DHV",
        },
        "extraction": {
            "method": "xlsx_qa_import",
            "text_source": "xlsx",
            "page_count": 1,
        },
        "pages": [{"page": 1, "text": text}],
        "sections": [{
            "section_id": "sec_001",
            "heading": "Câu hỏi và trả lời tuyển sinh 2026",
            "page_start": 1,
            "page_end": 1,
            "text": text,
        }],
        "records": records,
        "warnings": [],
    }
    return write_structured_json(document, output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", default=str(DEFAULT_INPUT))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    args = parser.parse_args()
    output = convert_workbook(args.input, args.output)
    print(f"Wrote {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())